from openpyxl import load_workbook

wb = load_workbook("lab_4xl.xlsx")
ws = wb["Sheet1"]
#file-d bicix
#a1-a5
ws.cell(row = 1, column = 1, value = "A")
ws.cell(row = 2, column = 1, value = 2)
ws.cell(row = 3, column = 1, value = 3)
ws.cell(row = 4, column = 1, value = 1)
ws.cell(row = 5, column = 1, value = 0)
# b1-b5
ws.cell(row = 1, column = 2, value = "B")
ws.cell(row = 2, column = 2, value = 5)
ws.cell(row = 3, column = 2, value = 7)
ws.cell(row = 4, column = 2, value = 7)
ws.cell(row = 5, column = 2, value = 3)
# c1-c5
ws.cell(row = 1, column = 3, value = "C")
ws.cell(row = 2, column = 3, value = 6)
ws.cell(row = 3, column = 3, value = 11)
ws.cell(row = 4, column = 3, value = 9)
ws.cell(row = 5, column = 3, value = 3)

# d1 - result
ws.cell(row = 1, column = 4, value = "Result")


for i in range(2,6):
  if int(ws.cell(row = i, column = 1).value) + int(ws.cell(row = i, column = 2).value) == int( ws.cell(row = i, column = 3).value):
    ws.cell(row = i, column = 4, value = "Зөв")
  else:
    ws.cell(row = i, column = 4, value = "Буруу")

# uurclultiig xadgalax
wb.save("lab_4xl.xlsx")


# file-s unshix
# print(ws['A2'].value)